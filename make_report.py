from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path('hull.kr_공개페이지_사이트분석보고서.docx')
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
for s in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[s]; st.font.name='Malgun Gothic'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
doc.styles['Normal'].font.size=Pt(9); doc.styles['Title'].font.size=Pt(25); doc.styles['Title'].font.bold=True; doc.styles['Title'].font.color.rgb=RGBColor(23,55,94)
doc.styles['Heading 1'].font.size=Pt(16); doc.styles['Heading 1'].font.color.rgb=RGBColor(31,78,121)
doc.styles['Heading 2'].font.size=Pt(12); doc.styles['Heading 2'].font.color.rgb=RGBColor(47,84,150)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def celltext(cell,text,bold=False,color=None):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold
    if color:r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],'1F4E79'); celltext(t.rows[0].cells[i],h,True,(255,255,255))
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): celltext(cells[i],v)
    doc.add_paragraph(''); return t
def shot(path, caption, width=5.7):
    if Path(path).exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path,width=Inches(width))
        cp=doc.add_paragraph(caption); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.runs[0].italic=True; cp.runs[0].font.size=Pt(8); cp.runs[0].font.color.rgb=RGBColor(100,100,100)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('hull.kr 공개 페이지 사이트 분석 보고서')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Playwright 기반 페이지 탐색·캡처 결과').bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('작성일: 2026년 9월 9일\n분석 범위: 로그인 없이 접근 가능한 공개 페이지').font.size=Pt(11)
shot('hull-home.png','그림 1. hull.kr 메인 화면(공개 상태)',5.8); doc.add_page_break()

doc.add_heading('1. 요약',1)
doc.add_paragraph('hull.kr은 한국어 기반의 개발 학습·기술 자료 공유 사이트로, 개발 노트와 커뮤니티를 중심으로 클라우드 네이티브, IoT/라즈베리파이, 우분투, 스프링 부트 관련 게시물을 제공한다. 상단 전역 메뉴와 게시판 분류가 일관되며, 일부 링크는 외부 시스템 또는 별도 서비스로 연결된다.')
table(['항목','확인 결과'],[('분석 일시','2026-09-09 (공개 페이지 기준)'),('접근 방식','Playwright로 페이지 이동, 본문 텍스트 확인, 전체 페이지 PNG 캡처'),('로그인 상태','제공된 인증정보로 로그인 시도했으나 실패: “가입된 회원아이디가 아니거나 비밀번호가 틀립니다.”'),('주요 콘텐츠','개발 노트, 강의/실습 게시판, 공지사항, 자료실, 질문답변, 새글 모음'),('외부/별도 서비스','NAS File Station, Spring Initializr, Oracle JDK 문서, JetBrains IntelliJ 다운로드'),('주의 관찰','최근 자유게시판에 1xbet 등 홍보성 게시물이 다수 노출됨')])
doc.add_heading('2. 분석 방법 및 범위',1)
doc.add_paragraph('다음 공개 URL을 순차 방문하고 페이지별 전체 화면을 캡처했다. 페이지의 제목, 핵심 본문 텍스트, 내부 링크 구조를 확인했으며, 로그인 이후 또는 권한 제한 콘텐츠는 분석 대상에서 제외했다.')
table(['구분','대상'],[('대표 페이지','홈, 개발 노트, 커뮤니티, 새글'),('게시판','클라우드네이티브, 스마트기기시스템, IoT 유무선 제어, 우분투, 스프링 부트, 공지사항, 자료실, 질문답변'),('도구/외부 연계','우분투 PPT, 스프링 부트 PPT, Spring Initializr, Java JDK 11, IntelliJ IDEA'),('산출물','페이지 캡처 17장 및 본 DOCX 보고서')])
doc.add_heading('3. 정보 구조 및 주요 기능',1)
table(['영역','기능 및 특징'],[('전역 헤더','회원가입, 로그인, 추가메뉴, 검색, 메뉴 버튼 제공'),('개발 노트','주제별 게시판을 묶어 최신 게시물과 분류별 글을 탐색'),('커뮤니티','IoT/도커·쿠버네티스 그룹별 게시물과 공지/자료/자유게시판 제공'),('게시판 기능','분류 탭, 게시물 수·페이지네이션, 작성자·작성일, 댓글 수 표시'),('자료/도구 링크','다운로드 또는 외부 서비스 페이지로 연결'),('새글 모음','전체 그룹, 검색 대상, 원글/댓글 필터와 검색어 입력 제공')])
doc.add_heading('4. 페이지별 분석 및 캡처',1)
items=[('홈','hull-page-01.png','핵심 게시판 미리보기와 커뮤니티 최신 글을 한 화면에 배치한다. 방문자가 학습 콘텐츠와 최근 활동을 빠르게 파악할 수 있다.'),('개발 노트','hull-page-02.png','IoT, 스마트기기시스템, 클라우드네이티브, 우분투, 스프링 부트 게시판을 통합 목록으로 보여준다.'),('클라우드네이티브','hull-page-03.png','공지/이슈/실습/이론 분류와 23건의 게시물 목록이 확인된다. 컨테이너·마이크로서비스 실습 중심이다.'),('스마트기기시스템','hull-page-04.png','22건의 게시물과 라즈베리파이·센서·MQTT 실습 콘텐츠가 중심이다.'),('IoT 유무선 제어','hull-page-05.png','15건의 게시물로, 센서 측정 및 GPIO/웹 제어 실습 흐름이 뚜렷하다.'),('우분투 리눅스','hull-page-06.png','20건의 설치·실습·이론 글이 있으며 VirtualBox, SSH, MariaDB, Apache/PHP 등 운영 주제를 다룬다.'),('우분투 PPT','hull-page-07.png','사이트 내부 게시판이 아닌 Synology NAS File Station 화면으로 연결된다.'),('스프링 부트','hull-page-08.png','52건의 글과 개발·이론·설정·설치·이슈 분류를 제공한다. Spring Security와 JPA 학습 글이 보인다.'),('스프링 부트 PPT','hull-page-09.png','Synology NAS File Station의 강의 교안 다운로드 폴더로 연결된다.'),('Spring Initializr','hull-page-10.png','프로젝트 생성 옵션과 의존성 추가 UI를 제공하는 외부형 도구 페이지다.'),('Java JDK 11','hull-page-11.png','Java SE 11 Reference Implementation 다운로드/설명 페이지로 연결된다.'),('IntelliJ IDEA','hull-page-12.png','JetBrains IntelliJ IDEA 다운로드 페이지로 연결되며 운영체제별 다운로드와 제품 안내를 제공한다.'),('커뮤니티','hull-page-13.png','IoT, 도커·쿠버네티스, 공지사항, 자료, 자유게시판을 그룹별로 묶은 허브다.'),('공지사항','hull-page-14.png','공지 1건이 표시되는 간결한 게시판이다.'),('자료실','hull-page-15.png','Java, YOLO, 설치 파일 등 코드·자료 중심의 20건 목록이 확인된다.'),('질문답변','hull-page-16.png','현재 공개 상태에서는 게시물이 없는 빈 게시판이다.'),('새글','hull-page-17.png','검색 조건과 최신 게시물 목록을 제공하며 여러 게시판의 활동을 통합 확인한다.')]
for i,(name,img,desc) in enumerate(items):
    doc.add_heading(f'{i+1}. {name}',2); doc.add_paragraph(desc); shot(img,f'그림 {i+2}. {name} 페이지 캡처')
    if i in [2,7,12]: doc.add_page_break()
doc.add_heading('5. UX·콘텐츠·운영 관점 평가',1)
table(['평가 영역','강점','개선 제안'],[('탐색성','주제별 게시판과 상단 메뉴가 명확함','메뉴에 게시물 수, 최신 업데이트, 인기 글을 보강'),('학습 콘텐츠','실습 중심의 연속적인 기술 자료가 풍부함','게시물 간 선수 지식/다음 글 링크와 태그 체계 추가'),('일관성','게시판 레이아웃과 공통 헤더가 반복되어 학습 비용이 낮음','외부 서비스 이동 시 새 탭·외부 링크 안내를 명시'),('검색','새글 모음에서 그룹/게시판/댓글 범위 검색을 지원','사이트 전체 통합 검색과 검색 결과 정렬·필터 개선'),('커뮤니티 품질','댓글 수와 작성자/날짜가 노출됨','자유게시판 스팸·홍보 글 탐지, 신고/승인, 신규 계정 제한 강화'),('접근성','본문 바로가기 링크와 입력 라벨이 존재함','아이콘-only 버튼의 대체 텍스트, 색 대비, 키보드 포커스 점검')])
doc.add_heading('6. 보안 및 신뢰성 관찰',1)
doc.add_paragraph('본 분석은 공개 화면 관찰에 한정되며 취약점 진단이나 침투 테스트가 아니다. 로그인 시도는 실패했으며, 비밀번호는 보고서에 기록하지 않았다. 공개 게시판에는 1xbet 관련 반복 홍보성 글과 무관한 영문/외국어 글이 다수 보여 콘텐츠 스팸 관리가 주요 운영 이슈로 판단된다. 또한 NAS File Station 등 별도 서비스로 연결되는 링크는 접근 권한, 세션 분리, 외부 노출 범위를 정기적으로 점검하는 것이 바람직하다.')
doc.add_heading('7. 우선순위별 권고사항',1)
table(['우선순위','권고사항','기대 효과'],[('높음','자유게시판 스팸 차단·신고·관리자 승인 정책 도입','콘텐츠 신뢰도와 커뮤니티 품질 개선'),('높음','외부/NAS 링크의 HTTPS·권한·세션·접근 로그 점검','자료 및 관리 화면의 노출 위험 감소'),('중간','통합 검색 및 주제 태그/관련 글 기능 강화','학습 자료 재발견성과 체류 시간 향상'),('중간','모바일·키보드·스크린리더 접근성 점검','다양한 이용 환경에서 사용성 개선'),('낮음','홈 화면에 카테고리 설명과 추천 학습 경로 추가','신규 방문자의 진입 장벽 완화')])
doc.add_heading('8. 결론',1)
doc.add_paragraph('hull.kr은 개발 교육과 실습 기록을 주제별 게시판으로 축적한 기술 커뮤니티형 사이트다. 핵심 정보 구조는 비교적 명료하고, IoT·리눅스·클라우드 네이티브·스프링 부트 자료가 서로 연결되어 있다는 점이 강점이다. 향후에는 공개 커뮤니티의 스팸 관리, 외부 서비스 링크의 신뢰성 안내, 검색/태그/학습 경로 개선을 우선 추진하면 콘텐츠 활용성과 운영 안정성을 함께 높일 수 있다.')
doc.save(out); print(out.resolve(), out.stat().st_size)
